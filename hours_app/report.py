"""Geração do relatório mensal de jornada em Excel (.xlsx).

Monta um arquivo com duas abas:
- "Horas": dias com registro no mês, com total em evidência, modalidade e
  feriados destacados com cor na linha.
- "Valores": tabela de constantes (config) + remuneração discriminada linha a
  linha + valor total da NF + chave PIX.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .constants import DAY_LABELS, WEEKDAY_ORDER
from .time_utils import minutes_to_duration_hhmm, parse_hhmm

CURRENCY_FORMAT = 'R$ #,##0.00'

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HOLIDAY_FILL = PatternFill("solid", fgColor="FCE4A6")
_TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
_SECTION_FILL = PatternFill("solid", fgColor="E2EFDA")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _lunch_duration_hhmm(row: pd.Series) -> str:
    """Duração do almoço em HH:MM (mesma lógica do export CSV)."""
    try:
        lunch_start = parse_hhmm(str(row["lunch_start_time"]))
        lunch_end = parse_hhmm(str(row["lunch_end_time"]))
        return minutes_to_duration_hhmm(max(lunch_end - lunch_start, 0))
    except Exception:
        return "00:00"


def _weekday_label(work_date: pd.Timestamp) -> str:
    return DAY_LABELS.get(WEEKDAY_ORDER[work_date.weekday()], "")


def _style_header(ws: Worksheet, row: int, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER


def _autosize(ws: Worksheet, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _build_horas_sheet(
    ws: Worksheet, df: pd.DataFrame, month_date: date
) -> None:
    headers = [
        "Dia",
        "Dia da semana",
        "Entrada",
        "Saída almoço",
        "Almoço",
        "Volta almoço",
        "Saída",
        "Modalidade",
        "Total (HH:MM)",
        "Feriado",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    month_start = month_date.replace(day=1)
    month_end = (pd.Timestamp(month_start) + pd.offsets.MonthEnd(0)).date()

    total_col = headers.index("Total (HH:MM)") + 1

    month_rows = pd.DataFrame()
    if not df.empty:
        in_month = (df["work_date"].dt.date >= month_start) & (
            df["work_date"].dt.date <= month_end
        )
        month_rows = df.loc[in_month].copy().sort_values(["work_date", "start_time"])

    total_minutes_month = 0
    for _, row in month_rows.iterrows():
        is_holiday = bool(row.get("is_holiday", False))
        minutes = 0 if is_holiday else int(row.get("total_minutes", 0) or 0)
        total_minutes_month += minutes
        work_mode = row.get("work_mode") or ""

        values = [
            row["work_date"].strftime("%d/%m/%Y"),
            _weekday_label(row["work_date"]),
            "" if is_holiday else str(row.get("start_time", "")),
            "" if is_holiday else str(row.get("lunch_start_time", "")),
            "" if is_holiday else _lunch_duration_hhmm(row),
            "" if is_holiday else str(row.get("lunch_end_time", "")),
            "" if is_holiday else str(row.get("end_time", "")),
            "Feriado" if is_holiday else str(work_mode),
            minutes_to_duration_hhmm(minutes),
            "Sim" if is_holiday else "Não",
        ]
        ws.append(values)
        excel_row = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=excel_row, column=col)
            cell.border = _BORDER
            cell.alignment = Alignment(horizontal="center")
            if is_holiday:
                cell.fill = _HOLIDAY_FILL
        # Total em evidência
        ws.cell(row=excel_row, column=total_col).font = Font(bold=True)

    # Linha de total do mês
    ws.append([])
    ws.append(["Total do mês", "", "", "", "", "", "", "",
               minutes_to_duration_hhmm(total_minutes_month), ""])
    total_row = ws.max_row
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=total_row, column=col)
        cell.fill = _TOTAL_FILL
        cell.font = Font(bold=True)
        cell.border = _BORDER
    ws.cell(row=total_row, column=1).alignment = Alignment(horizontal="left")

    _autosize(ws, [12, 14, 9, 12, 9, 12, 9, 14, 14, 9])
    ws.freeze_panes = "A2"


def _build_valores_sheet(
    ws: Worksheet, config: dict, breakdown: dict, pix_key: str
) -> None:
    bold = Font(bold=True)

    def section_title(text: str) -> None:
        ws.append([text])
        r = ws.max_row
        cell = ws.cell(row=r, column=1)
        cell.font = bold
        cell.fill = _SECTION_FILL

    def header_row(cols: list[str]) -> None:
        ws.append(cols)
        _style_header(ws, ws.max_row, len(cols))

    def money(row: int, col: int, value: float) -> None:
        cell = ws.cell(row=row, column=col, value=float(value))
        cell.number_format = CURRENCY_FORMAT

    # --- Tabela A: Constantes ---
    section_title("Constantes de remuneração")
    header_row(["Constante", "Valor", "Descrição"])
    constantes = [
        ("Remuneração base", config.get("valor_base", 0),
         "Valor fixo mensal"),
        ("Valor hora (ajuste)", config.get("valor_hora", 0),
         "R$ por hora trabalhada a mais ou a menos vs. a meta"),
        ("Auxílio transporte", config.get("valor_aux_transporte", 0),
         "R$ por dia trabalhado presencialmente"),
        ("Bonificação (meta)", config.get("valor_bonus", 0),
         "Aplicada quando a meta do mês é batida"),
    ]
    for nome, valor, descricao in constantes:
        ws.append([nome, None, descricao])
        money(ws.max_row, 2, valor)

    ws.append([])

    # --- Tabela B: Discriminação do mês ---
    section_title("Remuneração do mês (discriminada)")
    header_row(["Item", "Detalhe", "Valor"])

    bateu_meta = breakdown.get("bonus_total", 0) > 0
    if breakdown.get("horas_excedentes", 0) > 0:
        detalhe_horas = (
            f"{breakdown['horas_excedentes']:.2f}h extras × "
            f"R$ {float(config.get('valor_hora', 0)):.2f}"
        )
    elif breakdown.get("horas_faltantes", 0) > 0:
        detalhe_horas = (
            f"-{breakdown['horas_faltantes']:.2f}h faltantes × "
            f"R$ {float(config.get('valor_hora', 0)):.2f}"
        )
    else:
        detalhe_horas = "Sem ajuste (meta exata)"

    linhas = [
        ("Remuneração base", "", breakdown.get("valor_base", 0)),
        ("Ajuste de horas trabalhadas", detalhe_horas,
         breakdown.get("ajuste_horas", 0)),
        ("Bonificação", "Bateu a meta" if bateu_meta else "Não bateu a meta",
         breakdown.get("bonus_total", 0)),
        ("Dias presenciais",
         f"{breakdown.get('presencial_days', 0)} dia(s) × "
         f"R$ {float(config.get('valor_aux_transporte', 0)):.2f}",
         breakdown.get("aux_transporte_total", 0)),
    ]
    for item, detalhe, valor in linhas:
        ws.append([item, detalhe, None])
        money(ws.max_row, 3, valor)

    # Total da remuneração
    ws.append(["Total da remuneração", "", None])
    total_row = ws.max_row
    money(total_row, 3, breakdown.get("total_remuneracao", 0))
    for col in range(1, 4):
        cell = ws.cell(row=total_row, column=col)
        cell.font = bold
        cell.fill = _TOTAL_FILL

    ws.append([])

    # --- Nota fiscal ---
    section_title("Nota fiscal")
    ws.append(["Valor total da NF", "", None])
    nf_row = ws.max_row
    money(nf_row, 3, breakdown.get("total_remuneracao", 0))
    for col in range(1, 4):
        ws.cell(row=nf_row, column=col).font = bold
    ws.append(["Chave PIX", pix_key or "", ""])
    ws.cell(row=ws.max_row, column=1).font = bold

    _autosize(ws, [28, 40, 16])


def build_monthly_report_xlsx(
    df: pd.DataFrame,
    month_date: date,
    config: dict,
    breakdown: dict,
    pix_key: str,
    day_target_minutes: int,
) -> bytes:
    """Monta o relatório mensal e retorna os bytes do arquivo .xlsx."""
    wb = Workbook()
    ws_horas = wb.active
    ws_horas.title = "Horas"
    _build_horas_sheet(ws_horas, df, month_date)

    ws_valores = wb.create_sheet("Valores")
    _build_valores_sheet(ws_valores, config, breakdown, pix_key)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
